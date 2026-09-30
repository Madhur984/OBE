"""
Word and Excel writers for the SAR annexures.

`sar_export.py` produces the tables as data. This module renders them into
the two formats a college actually submits and works in: a .docx whose tables
can be pasted straight into the SAR without reformatting, and an .xlsx that
keeps the working visible so a peer team can trace a figure.

The distinction matters more than it sounds. Markdown and CSV are fine for
this codebase's own tests, but nobody submits a markdown table to NBA, and
handing faculty a CSV means they retype it into Word -- which is where
transcription errors enter a document nobody re-checks.

Both writers are optional-dependency-guarded. A bare install still computes
every number; it just cannot render them, and says so rather than failing at
the point of export.
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from app.engine.sar_export import Table

try:  # python-docx
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    _DOCX = True
except Exception:                                    # pragma: no cover
    _DOCX = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    _XLSX = True
except Exception:                                    # pragma: no cover
    _XLSX = False


class ExportUnavailable(RuntimeError):
    """The optional renderer for this format is not installed."""


def formats_available() -> dict[str, bool]:
    return {"docx": _DOCX, "xlsx": _XLSX}


@dataclass(frozen=True)
class DocumentMeta:
    institution: str = ""
    program: str = ""
    framework: str = ""
    academic_year: str = ""
    prepared_by: str = ""
    note: str = ""


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

_HEADER_BG = "D9E2F3"


def _shade(cell, hex_fill: str) -> None:
    """python-docx has no shading API, so the XML is written directly."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _add_table(doc, table: Table, *, font_pt: float = 8.5) -> None:
    doc.add_paragraph()
    cap = doc.add_paragraph()
    run = cap.add_run(f"Table No. {table.number}: ")
    run.bold = True
    cap.add_run(table.title)

    t = doc.add_table(rows=1, cols=len(table.headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(table.headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(font_pt)
        _shade(cell, _HEADER_BG)

    for row in table.rows:
        cells = t.add_row().cells
        # The SAR's own convention: the trailing average row is emphasised,
        # because it is the figure the peer team reads first.
        is_summary = bool(row) and row[0].lower().startswith(
            ("direct attainment", "average", "overall")
        )
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                           else WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(val)
            r.font.size = Pt(font_pt)
            r.bold = is_summary
            # A '-' is not a zero; grey it so the eye does not read it as one.
            if val == "-":
                r.font.color.rgb = RGBColor(0x90, 0x90, 0x90)


DEFAULT_TITLE = "Outcome attainment - SAR annexures"


def tables_to_docx(tables: list[Table], *, meta: DocumentMeta | None = None,
                   notes: list[str] | None = None,
                   title: str = DEFAULT_TITLE) -> bytes:
    """Render SAR tables into a Word document, returned as bytes."""
    if not _DOCX:
        raise ExportUnavailable(
            "Word export needs python-docx. Install it with "
            "`pip install python-docx`; every number is still computed without it."
        )
    m = meta or DocumentMeta()
    doc = Document()

    doc.add_heading(title, level=1)
    if any([m.institution, m.program, m.framework, m.academic_year, m.prepared_by]):
        info = doc.add_table(rows=0, cols=2)
        info.style = "Table Grid"
        for label, value in (
            ("Institution", m.institution), ("Program", m.program),
            ("Outcome framework", m.framework), ("Academic year", m.academic_year),
            ("Prepared by", m.prepared_by),
        ):
            if not value:
                continue
            cells = info.add_row().cells
            cells[0].text = ""
            r = cells[0].paragraphs[0].add_run(label)
            r.bold = True
            cells[1].text = value
    if m.note:
        doc.add_paragraph(m.note)

    for table in tables:
        _add_table(doc, table)

    if notes:
        doc.add_heading("Methodology notes", level=2)
        for n in notes:
            doc.add_paragraph(n, style="List Bullet")

    # The one caveat every generated SAR annexure should carry.
    p = doc.add_paragraph()
    r = p.add_run(
        "A dash indicates the outcome is not addressed by that course. It is "
        "distinct from a zero, which would mean addressed and not attained."
    )
    r.italic = True
    r.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_THIN = None


def _borders():
    global _THIN
    if _THIN is None:
        side = Side(style="thin", color="B0B0B0")
        _THIN = Border(left=side, right=side, top=side, bottom=side)
    return _THIN


def _sheet_name(table: Table) -> str:
    # Excel sheet names cap at 31 chars and reject : \ / ? * [ ]
    raw = f"{table.number} {table.title}"
    for ch in ':\\/?*[]':
        raw = raw.replace(ch, "-")
    return raw[:31]


def _write_cover(ws, *, meta: DocumentMeta, notes: list[str] | None,
                 contents: list[tuple[str, str]],
                 heading: str = "Outcome attainment — SAR annexures") -> None:
    """The cover sheet, shared by both workbook writers."""
    ws["A1"] = heading
    ws["A1"].font = Font(bold=True, size=14)
    row = 3
    for label, value in (
        ("Institution", meta.institution), ("Program", meta.program),
        ("Outcome framework", meta.framework),
        ("Academic year", meta.academic_year), ("Prepared by", meta.prepared_by),
    ):
        if not value:
            continue
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Contents").font = Font(bold=True)
    row += 1
    for left, right in contents:
        ws.cell(row=row, column=1, value=left)
        ws.cell(row=row, column=2, value=right)
        row += 1
    if notes:
        row += 1
        ws.cell(row=row, column=1, value="Methodology notes").font = Font(bold=True)
        row += 1
        for n in notes:
            c = ws.cell(row=row, column=1, value=n)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            row += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70


def tables_to_xlsx(tables: list[Table], *, meta: DocumentMeta | None = None,
                   notes: list[str] | None = None,
                   title: str = DEFAULT_TITLE) -> bytes:
    """
    Render SAR tables into a workbook, one sheet per table plus a cover sheet.

    Numbers are written as numbers, not strings, so the college can check the
    arithmetic in the spreadsheet they already trust. A '-' stays text,
    because it is not a value.
    """
    if not _XLSX:
        raise ExportUnavailable(
            "Excel export needs openpyxl. Install it with `pip install openpyxl`; "
            "every number is still computed without it."
        )
    m = meta or DocumentMeta()
    wb = Workbook()

    cover = wb.active
    cover.title = "Cover"
    _write_cover(cover, meta=m, notes=notes,
                 contents=[(f"Table {t.number}", t.title) for t in tables])

    head_fill = PatternFill("solid", fgColor=_HEADER_BG)
    for table in tables:
        ws = wb.create_sheet(_sheet_name(table))
        ws.cell(row=1, column=1, value=f"Table No. {table.number}: {table.title}").font = \
            Font(bold=True, size=11)

        for j, h in enumerate(table.headers, start=1):
            c = ws.cell(row=3, column=j, value=h)
            c.font = Font(bold=True, size=9)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = _borders()

        for i, row_vals in enumerate(table.rows, start=4):
            is_summary = bool(row_vals) and row_vals[0].lower().startswith(
                ("direct attainment", "average", "overall")
            )
            for j, val in enumerate(row_vals, start=1):
                value: object = val
                if j > 1 and val not in ("-", ""):
                    try:
                        value = float(val)
                    except ValueError:
                        value = val
                c = ws.cell(row=i, column=j, value=value)
                c.font = Font(size=9, bold=is_summary,
                              color="909090" if val == "-" else "000000")
                c.alignment = Alignment(horizontal="left" if j == 1 else "center")
                c.border = _borders()
                if isinstance(value, float):
                    c.number_format = "0.00"

        ws.column_dimensions["A"].width = 34
        if len(table.headers) > 1 and table.headers[1].lower().startswith("course name"):
            ws.column_dimensions["B"].width = 40
            start = 3
        else:
            start = 2
        for j in range(start, len(table.headers) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 8
        ws.freeze_panes = ws.cell(row=4, column=2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def course_blocks_to_xlsx(sheets: dict[str, list[Table]], *,
                          meta: DocumentMeta | None = None,
                          notes: list[str] | None = None,
                          title: str = "CO - PO / PSO mapping",
                          label_width: int = 62, col_width: int = 6,
                          spacer_rows: int = 1) -> bytes:
    """
    Stack many course blocks down one worksheet, the way a department's own
    mapping workbook is laid out.

    `sheets` maps a worksheet name to its ordered blocks, so a whole programme
    goes out as "2nd year" / "3rd year" / "4th year" in one file. This is a
    different shape from `tables_to_xlsx`, which puts one table on one sheet
    with the header pinned at row 3; both are wanted, so neither is bent to
    fit the other.

    Two things this writer does that `tables_to_xlsx` cannot. A mapping level
    is written as the integer 3, not the float 3.00 -- levels are 1/2/3 and
    printing 3.00 in a matrix cell is wrong, while the average row genuinely
    is a decimal. And an empty string is written as a genuinely empty cell
    rather than being coerced, because in this layout blank means "not
    mapped" and 0 would read as a scored zero.
    """
    if not _XLSX:
        raise ExportUnavailable(
            "Excel export needs openpyxl. Install it with `pip install openpyxl`; "
            "every number is still computed without it."
        )
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    _write_cover(
        cover, meta=meta or DocumentMeta(), notes=notes,
        contents=[(name, f"{len(blocks)} course block(s)")
                  for name, blocks in sheets.items()],
        heading=title,
    )

    head_fill = PatternFill("solid", fgColor=_HEADER_BG)
    used: set[str] = {"Cover"}
    for sheet_name, blocks in sheets.items():
        name = _unique_sheet_name(sheet_name or "Mapping", used)
        used.add(name)
        ws = wb.create_sheet(name)
        row = 1
        widest = 1
        for table in blocks:
            widest = max(widest, len(table.headers))
            for j, h in enumerate(table.headers, start=1):
                c = ws.cell(row=row, column=j, value=h)
                c.font = Font(bold=True, size=9)
                c.fill = head_fill
                c.border = _borders()
                c.alignment = Alignment(
                    horizontal="left" if j == 1 else "center",
                    vertical="center", wrap_text=True)
            row += 1

            for i, row_vals in enumerate(table.rows):
                is_summary = i in table.summary_rows
                for j, val in enumerate(row_vals, start=1):
                    c = ws.cell(row=row, column=j, value=_cell_value(val, j, is_summary))
                    c.font = Font(size=9, bold=is_summary)
                    c.border = _borders()
                    c.alignment = (
                        Alignment(wrap_text=True, vertical="top", horizontal="left")
                        if j == 1 else
                        Alignment(horizontal="center", vertical="center"))
                    # No number format on the average: the string was already
                    # rounded to two places, and General renders it the way
                    # the department's own workbook does -- 3, not 3.00.
                    if isinstance(c.value, int) and not isinstance(c.value, bool):
                        c.number_format = "0"
                row += 1
            row += spacer_rows

        ws.column_dimensions["A"].width = label_width
        for j in range(2, widest + 1):
            ws.column_dimensions[get_column_letter(j)].width = col_width
        # Only the label column. Freezing rows as well would pin whichever
        # course happened to be first, which is meaningless on a stacked sheet.
        ws.freeze_panes = "B1"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_value(val: str, column: int, is_summary: bool) -> object:
    """A level is an int, an average is a float, and blank stays blank."""
    if column == 1:
        return val
    if val in ("", "-"):
        return None if val == "" else val
    try:
        return float(val) if is_summary else int(val)
    except ValueError:
        try:
            return float(val)
        except ValueError:
            return val


def _unique_sheet_name(raw: str, used: set[str]) -> str:
    base = raw
    for ch in ':\\/?*[]':
        base = base.replace(ch, "-")
    base = base[:31] or "Sheet"
    if base not in used:
        return base
    for n in range(2, 100):
        suffix = f" {n}"
        candidate = base[:31 - len(suffix)] + suffix
        if candidate not in used:
            return candidate
    return base[:28] + "~"


def evidence_bundle_to_docx(bundle_markdown: str, *, title: str = "Evidence bundle") -> bytes:
    """
    Render an evidence bundle's markdown into Word.

    A deliberately small markdown subset -- headings, tables, paragraphs --
    because the input is generated by `audit.EvidenceBundle.to_markdown`, not
    written by a human, so the full grammar is not in play.
    """
    if not _DOCX:
        raise ExportUnavailable("Word export needs python-docx.")
    doc = Document()
    pending: list[list[str]] = []

    def flush() -> None:
        if not pending:
            return
        header, *body = pending
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(8.5)
            _shade(cell, _HEADER_BG)
        for row in body:
            cells = t.add_row().cells
            for i, val in enumerate(row[:len(header)]):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(val)
                r.font.size = Pt(8.5)
        pending.clear()
        doc.add_paragraph()

    for line in bundle_markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(set(c) <= set("-: ") and c for c in cells):
                continue                      # the ---|--- separator row
            pending.append(cells)
            continue
        flush()
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("```"):
            continue
        elif stripped:
            doc.add_paragraph(stripped.replace("**", ""))
    flush()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
